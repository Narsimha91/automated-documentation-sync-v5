from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("README_Auto_Sync_SDLC.pptx")
WIDTH = 13.333
HEIGHT = 7.5

INK = "173247"
NAVY = "102A43"
BLUE = "276678"
TEAL = "168C8C"
MINT = "DDF3EF"
CORAL = "D66B54"
PALE_CORAL = "F9E8E2"
GOLD = "E9B44C"
PAPER = "F4F7F8"
WHITE = "FFFFFF"
MUTED = "637987"
LINE = "D9E2E7"
GREEN = "2D805A"
PALE_GREEN = "E4F3EA"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, width, height, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1)
    return shape


def text(
    slide,
    x,
    y,
    width,
    height,
    value,
    size=16,
    fill=INK,
    bold=False,
    font="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    italic=False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(fill)
    return box


def line(slide, x1, y1, x2, y2, fill=LINE, width=1.5):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color(fill)
    connector.line.width = Pt(width)
    return connector


def new_slide(prs, background=PAPER):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(background)
    return slide


def footer(slide, page, dark=False):
    fill = "B3C5CE" if dark else MUTED
    line(slide, 0.6, 7.08, 12.73, 7.08, "315067" if dark else LINE, 0.8)
    text(slide, 0.65, 7.15, 8.5, 0.2, "README AUTO-SYNC  /  SDLC CASE STUDY", 8, fill, bold=True)
    text(slide, 12.0, 7.13, 0.65, 0.22, f"{page:02d}", 9, fill, bold=True, align=PP_ALIGN.RIGHT)


def heading(slide, number, title_value, subtitle=None):
    text(slide, 0.7, 0.42, 1.2, 0.22, number.upper(), 9, TEAL, bold=True)
    text(slide, 0.7, 0.75, 11.9, 0.55, title_value, 27, NAVY, bold=True, font="Aptos Display")
    if subtitle:
        text(slide, 0.72, 1.38, 11.8, 0.42, subtitle, 12, MUTED)


def card(slide, x, y, width, height, title_value, body, accent=TEAL, bg=WHITE, title_size=15, body_size=11):
    rect(slide, x, y, width, height, bg)
    rect(slide, x, y, 0.07, height, accent, radius=False)
    text(slide, x + 0.22, y + 0.18, width - 0.42, 0.36, title_value, title_size, NAVY, bold=True)
    text(slide, x + 0.22, y + 0.62, width - 0.42, height - 0.76, body, body_size, MUTED)


def add_bullet(slide, x, y, width, value, marker_color=TEAL, size=13):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + 0.08), Inches(0.1), Inches(0.1))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color(marker_color)
    dot.line.fill.background()
    text(slide, x + 0.22, y, width - 0.22, 0.38, value, size, INK)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    prs.core_properties.title = "README Auto-Sync: End-to-End SDLC"
    prs.core_properties.subject = "Human-gated SDLC pipeline and GitHub Actions README synchronization"
    prs.core_properties.author = "GitHub Copilot"
    prs.core_properties.keywords = "SDLC, GitHub Actions, README, human-in-the-loop"

    # 1. Cover
    slide = new_slide(prs, NAVY)
    rect(slide, 0.74, 0.78, 0.12, 5.45, TEAL, radius=False)
    text(slide, 1.12, 0.82, 6.5, 0.34, "AUTOMATED DOCUMENTATION  /  CASE STUDY", 10, "91D5D0", bold=True)
    text(slide, 1.08, 1.48, 7.2, 1.4, "README\nAUTO-SYNC", 34, WHITE, bold=True, font="Aptos Display")
    text(slide, 1.12, 3.28, 6.1, 0.9, "From Confluence story to a human-approved,\nworking documentation pipeline", 19, "D7E5EA")
    text(slide, 1.12, 5.83, 5.9, 0.32, "SDLC WALKTHROUGH  •  07 OCT 2026", 10, "AFC0C9", bold=True)
    # Diagram motif is a workflow, not decoration.
    steps = [("STORY", GOLD), ("BUILD", "4FB7AE"), ("VERIFY", "8CCB8B"), ("SHIP", "F09072")]
    x_positions = [9.1, 10.1, 11.1, 12.1]
    line(slide, 9.1, 3.08, 12.25, 3.08, "527183", 2)
    for (label, fill), x in zip(steps, x_positions):
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(2.68), Inches(0.52), Inches(0.52))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color(fill)
        circle.line.fill.background()
        text(slide, x - 0.18, 3.38, 0.9, 0.25, label, 8, "D7E5EA", bold=True, align=PP_ALIGN.CENTER)
    text(slide, 9.05, 4.18, 3.35, 1.15, "A small change,\nwith a complete trail", 18, WHITE, bold=True, font="Aptos Display")
    footer(slide, 1, dark=True)

    # 2. User story and contract
    slide = new_slide(prs)
    heading(slide, "01  /  The brief", "Keep the README feature list in sync", "The user story defines a narrow automation boundary and a fail-fast publishing path.")
    rect(slide, 0.72, 1.95, 11.9, 1.28, NAVY)
    text(slide, 1.0, 2.14, 11.3, 0.85, "“As a developer, I want an automated GitHub Actions workflow to update our README features list whenever src/features.py changes on main, so documentation stays synced without manual effort.”", 16, WHITE, bold=True)
    items = [
        ("TRIGGER", "Only a push to main\nthat changes src/features.py", TEAL),
        ("TEST GATE", "Run pytest first; any\nfailure stops publication", GREEN),
        ("SYNC REGION", "Validate one ordered\nstart/end marker pair", BLUE),
        ("PRESERVE", "Never rewrite content\noutside the markers", CORAL),
        ("PUBLISH", "Create or update a PR\nfrom docs/auto-update-readme", GOLD),
    ]
    x0, gap, w = 0.72, 0.18, 2.28
    for index, (label, body, accent) in enumerate(items):
        x = x0 + index * (w + gap)
        rect(slide, x, 3.63, w, 2.2, WHITE)
        rect(slide, x, 3.63, w, 0.08, accent, radius=False)
        text(slide, x + 0.18, 3.91, w - 0.36, 0.28, label, 10, accent, bold=True)
        text(slide, x + 0.18, 4.37, w - 0.36, 1.15, body, 13, NAVY, bold=True)
    text(slide, 0.78, 6.2, 11.7, 0.35, "Design choice: feature names come from features() at runtime; Create, Update, and READ are data, not workflow constants.", 11, MUTED)
    footer(slide, 2)

    # 3. Gated SDLC
    slide = new_slide(prs)
    heading(slide, "02  /  Governance", "Eight phases, explicit human gates", "Artifacts moved forward only after the user reviewed and approved each phase.")
    phases = [
        ("01", "Requirements", "Story distilled;\nformat + marker rules agreed", TEAL),
        ("02", "Architecture", "Python 3.12, pytest,\nsync script, one workflow", BLUE),
        ("03", "Design review", "No blockers; four\nsafeguards approved", GREEN),
        ("04", "Implementation plan", "Ordered tasks +\npermissions gate approved", GOLD),
        ("05", "Implementation", "Feature module, sync,\ntests, Actions workflow", CORAL),
        ("06", "Code review", "No blockers; pytest\npinning deferred", BLUE),
        ("07", "Verification", "11 tests passed on\nPython 3.12.10", GREEN),
        ("08", "Pull requests", "Human-approved PRs;\nmerge triggers automation", TEAL),
    ]
    for i, (number, name, detail, accent) in enumerate(phases):
        row, col = divmod(i, 4)
        x = 0.72 + col * 3.04
        y = 2.0 + row * 2.13
        rect(slide, x, y, 2.82, 1.72, WHITE)
        rect(slide, x, y, 0.62, 0.62, accent)
        text(slide, x, y + 0.12, 0.62, 0.28, number, 13, WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(slide, x + 0.8, y + 0.13, 1.85, 0.48, name, 14, NAVY, bold=True)
        text(slide, x + 0.18, y + 0.82, 2.48, 0.66, detail, 11, MUTED)
        text(slide, x + 0.18, y + 1.5, 2.4, 0.17, "HUMAN APPROVAL GATE", 7, accent, bold=True)
    footer(slide, 3)

    # 4. Architecture flow
    slide = new_slide(prs)
    heading(slide, "03  /  Architecture", "The pipeline is intentionally linear", "A failed test or invalid marker stops before README write or PR publication.")
    flow = [
        ("PUSH", "main +\nsrc/features.py", BLUE),
        ("TEST", "pytest suite\nfirst", GREEN),
        ("COMPUTE", "features() →\nordered bullets", TEAL),
        ("VALIDATE", "unique, ordered\nmarker pair", GOLD),
        ("UPDATE", "replace marker\ninterior only", CORAL),
        ("PUBLISH", "bot branch +\ncreate/update PR", NAVY),
    ]
    left, y, w, h, gap = 0.67, 2.42, 1.78, 1.45, 0.32
    mid_y = y + h / 2
    for i in range(len(flow) - 1):
        x1 = left + i * (w + gap) + w
        line(slide, x1 + 0.02, mid_y, x1 + gap - 0.02, mid_y, "AABCC5", 2)
    for i, (label, body, accent) in enumerate(flow):
        x = left + i * (w + gap)
        rect(slide, x, y, w, h, WHITE)
        rect(slide, x, y, w, 0.1, accent, radius=False)
        text(slide, x + 0.13, y + 0.27, w - 0.26, 0.27, label, 9, accent, bold=True)
        text(slide, x + 0.13, y + 0.66, w - 0.26, 0.64, body, 12, NAVY, bold=True)
    rect(slide, 0.72, 4.48, 5.65, 1.22, PALE_CORAL)
    text(slide, 0.95, 4.7, 1.08, 0.25, "STOP", 10, CORAL, bold=True)
    text(slide, 2.0, 4.66, 4.05, 0.66, "Tests fail OR markers invalid → no README write; publishing step is never reached.", 12, INK, bold=True)
    rect(slide, 6.65, 4.48, 5.95, 1.22, MINT)
    text(slide, 6.9, 4.7, 1.32, 0.25, "PUBLISH", 10, TEAL, bold=True)
    text(slide, 8.18, 4.66, 4.05, 0.66, "Only a real README diff updates docs/auto-update-readme and its PR to main.", 12, INK, bold=True)
    text(slide, 0.8, 6.12, 11.8, 0.35, "Permissions: contents: write + pull-requests: write; GITHUB_TOKEN; actions pinned to immutable SHAs.", 11, MUTED)
    footer(slide, 4)

    # 5. Safety contract
    slide = new_slide(prs)
    heading(slide, "04  /  Safety", "The markers are a hard ownership boundary", "Automation owns the interior only; the rest of the document remains human-maintained.")
    rect(slide, 0.72, 2.0, 5.4, 3.95, NAVY)
    text(slide, 1.0, 2.28, 4.8, 0.32, "README.md  /  managed region", 11, "91D5D0", bold=True)
    code_lines = [
        ("<!-- docs-sync: start -->", "80D2C8"),
        ("- Create", WHITE),
        ("- Update", WHITE),
        ("- READ", WHITE),
        ("<!-- docs-sync: end -->", "80D2C8"),
    ]
    for index, (value, fill) in enumerate(code_lines):
        text(slide, 1.05, 2.95 + index * 0.48, 4.7, 0.3, value, 15, fill, font="Consolas", bold=(index in (0, 4)))
    text(slide, 1.03, 5.55, 4.72, 0.24, "Outside these markers: untouched", 10, "B9CBD4", italic=True)
    rules = [
        ("01", "Exactly one start + one end", "Missing, duplicate, or reversed markers fail before write."),
        ("02", "Stable rendering", "One ordered Markdown bullet per feature; embedded line breaks become spaces."),
        ("03", "Exact preservation", "Tests compare surrounding bytes, including CRLF context."),
        ("04", "Idempotent run", "If generated content matches, do not write or create an empty commit."),
    ]
    for index, (number, title_value, body) in enumerate(rules):
        y = 2.02 + index * 0.98
        text(slide, 6.55, y, 0.42, 0.26, number, 10, TEAL, bold=True)
        text(slide, 7.08, y - 0.03, 5.1, 0.28, title_value, 13, NAVY, bold=True)
        text(slide, 7.08, y + 0.31, 5.12, 0.49, body, 10, MUTED)
    footer(slide, 5)

    # 6. Implementation map
    slide = new_slide(prs)
    heading(slide, "05  /  Build", "Small components, clear ownership", "Keep business data in Python, sync logic testable, and YAML focused on orchestration.")
    components = [
        ("src/features.py", "Source of truth", "FEATURES = [Create, Update, READ]\nfeatures() returns a copy.", TEAL),
        ("scripts/sync_readme.py", "Sync behavior", "Validate markers, render bullets, compare output, and write only on change.", BLUE),
        ("tests/", "Quality gate", "Exact feature values plus rendering, marker, preservation, error, and no-op cases.", GREEN),
        (".github/workflows/", "Automation", "main + path filter → Python 3.12 → pytest → sync → PR action.", CORAL),
    ]
    for i, (path, title_value, body, accent) in enumerate(components):
        col = i % 2
        row = i // 2
        x, y = 0.72 + col * 6.08, 2.03 + row * 1.75
        rect(slide, x, y, 5.72, 1.48, WHITE)
        text(slide, x + 0.22, y + 0.16, 5.2, 0.22, path, 10, accent, bold=True, font="Consolas")
        text(slide, x + 0.22, y + 0.48, 5.2, 0.28, title_value, 15, NAVY, bold=True)
        text(slide, x + 0.22, y + 0.86, 5.2, 0.5, body, 10, MUTED)
    rect(slide, 0.72, 5.82, 11.8, 0.75, MINT)
    text(slide, 0.98, 6.04, 11.3, 0.3, "Security and reproducibility: least write permissions • GITHUB_TOKEN PR creation enabled • actions pinned to commit SHAs", 11, INK, bold=True)
    footer(slide, 6)

    # 7. Review and verification
    slide = new_slide(prs)
    heading(slide, "06  /  Assurance", "Evidence from review and runtime", "The code was checked locally, then the real main-branch trigger exercised the workflow twice.")
    rect(slide, 0.72, 2.0, 3.48, 3.95, NAVY)
    text(slide, 1.0, 2.34, 2.9, 0.3, "PYTHON 3.12.10", 10, "91D5D0", bold=True)
    text(slide, 0.98, 2.92, 2.95, 0.9, "11 / 11", 31, WHITE, bold=True, font="Aptos Display")
    text(slide, 1.0, 3.88, 2.85, 0.65, "pytest tests passed", 16, WHITE, bold=True)
    text(slide, 1.0, 4.82, 2.85, 0.8, "Disposable README check: bullets correct, outside bytes preserved, second run unchanged.", 11, "D7E5EA")
    card(slide, 4.55, 2.0, 3.72, 1.63, "CODE REVIEW", "No blocking findings. Pytest pinning was a low-severity recommendation and was explicitly deferred.", GREEN, title_size=13, body_size=10)
    card(slide, 8.55, 2.0, 4.05, 1.63, "WORKFLOW RUN #1", "Successful after PR #1 merged; opened README PR #2 for Create + Update.", TEAL, title_size=13, body_size=10)
    card(slide, 4.55, 4.04, 3.72, 1.63, "WORKFLOW RUN #2", "Successful after PR #3 merged; opened README PR #4 with READ added.", BLUE, title_size=13, body_size=10)
    card(slide, 8.55, 4.04, 4.05, 1.63, "BOUNDARY", "The human-approved source PR and bot-generated README PR were separately reviewed and merged.", CORAL, title_size=13, body_size=10)
    text(slide, 0.78, 6.28, 11.7, 0.3, "Action runs were observed on GitHub; local tests alone were not treated as end-to-end workflow evidence.", 10, MUTED)
    footer(slide, 7)

    # 8. Actual PR timeline
    slide = new_slide(prs)
    heading(slide, "07  /  Execution", "Two change cycles proved the handoff", "Source-code PRs changed the feature list; the workflow generated a separate README PR each time.")
    timeline = [
        ("PR #1", "Pipeline foundation", "Merged implementation: source, sync script, tests, workflow.", BLUE),
        ("PR #2", "Initial README sync", "Bot PR added Create + Update. Reviewed and merged.", TEAL),
        ("PR #3", "Add READ", "Feature source and exact-value test. 11 tests passed; merged.", CORAL),
        ("PR #4", "README sync for READ", "Bot PR added - READ only. Reviewed and merged.", GREEN),
    ]
    line(slide, 1.27, 3.05, 12.1, 3.05, "AABCC5", 2)
    for i, (label, title_value, body, accent) in enumerate(timeline):
        x = 0.72 + i * 3.04
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.16), Inches(2.78), Inches(0.54), Inches(0.54))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color(accent)
        circle.line.fill.background()
        rect(slide, x, 3.55, 2.76, 1.96, WHITE)
        text(slide, x + 0.18, 3.77, 2.35, 0.25, label, 10, accent, bold=True)
        text(slide, x + 0.18, 4.14, 2.4, 0.4, title_value, 14, NAVY, bold=True)
        text(slide, x + 0.18, 4.68, 2.4, 0.64, body, 10, MUTED)
    rect(slide, 0.72, 5.94, 11.8, 0.67, PALE_GREEN)
    text(slide, 0.97, 6.14, 11.2, 0.25, "Outcome: the source change was merged, the workflow ran, and its generated README diff matched the feature list.", 11, GREEN, bold=True)
    footer(slide, 8)

    # 9. Repeatable operating model
    slide = new_slide(prs)
    heading(slide, "08  /  How to use it", "The next feature follows the same route", "This separation keeps the source PR reviewable and makes the README change independently visible.")
    cycle = [
        ("1", "Edit source", "Add a name to\nsrc/features.py", TEAL),
        ("2", "Test + PR", "Run pytest; open a\nsource PR to main", BLUE),
        ("3", "Merge source", "Push to main matches\nthe workflow trigger", CORAL),
        ("4", "Bot sync", "Tests pass; marker\nregion is regenerated", GOLD),
        ("5", "Review README PR", "Bot opens/updates\ndocs/auto-update-readme", GREEN),
    ]
    x0, w, gap, y = 0.72, 2.16, 0.25, 2.55
    for i, (number, title_value, body, accent) in enumerate(cycle):
        x = x0 + i * (w + gap)
        rect(slide, x, y, w, 2.08, WHITE)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18), Inches(y + 0.18), Inches(0.48), Inches(0.48))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color(accent)
        circle.line.fill.background()
        text(slide, x + 0.18, y + 0.26, 0.48, 0.2, number, 11, WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(slide, x + 0.18, y + 0.86, w - 0.35, 0.38, title_value, 14, NAVY, bold=True)
        text(slide, x + 0.18, y + 1.38, w - 0.35, 0.54, body, 10, MUTED)
        if i < len(cycle) - 1:
            line(slide, x + w, y + 1.02, x + w + gap - 0.03, y + 1.02, "AABCC5", 1.5)
    rect(slide, 0.72, 5.35, 11.78, 1.12, NAVY)
    text(slide, 1.0, 5.58, 2.0, 0.3, "HUMAN CONTROL", 10, "91D5D0", bold=True)
    text(slide, 3.05, 5.53, 8.95, 0.58, "Review and merge the source PR; separately review and merge the bot README PR. The workflow never writes directly to main.", 13, WHITE, bold=True)
    footer(slide, 9)

    # 10. Closing and references
    slide = new_slide(prs, NAVY)
    text(slide, 0.82, 0.7, 1.15, 0.25, "09  /  RESULT", 9, "91D5D0", bold=True)
    text(slide, 0.8, 1.15, 11.7, 0.86, "A small automation,\nwith a visible audit trail.", 29, WHITE, bold=True, font="Aptos Display")
    highlights = [
        ("3", "features in README", TEAL),
        ("11/11", "tests passed", GREEN),
        ("2", "successful workflow runs", GOLD),
        ("4", "PRs merged", CORAL),
    ]
    for i, (value, label, accent) in enumerate(highlights):
        x = 0.82 + i * 3.08
        rect(slide, x, 3.18, 2.77, 1.34, "1C3B53")
        text(slide, x + 0.2, 3.39, 2.35, 0.52, value, 24, accent, bold=True, font="Aptos Display")
        text(slide, x + 0.2, 4.02, 2.35, 0.26, label, 10, "D7E5EA", bold=True)
    text(slide, 0.86, 5.1, 11.55, 0.62, "Presentation source: docs/nsg-agent-conversation.md\nWorkflow: .github/workflows/readme-auto-sync.yml  •  Tests: tests/  •  Verification: tests/results/test_results.md", 10, "C1D0D7")
    text(slide, 0.86, 6.23, 11.4, 0.3, "README AUTO-SYNC  /  END-TO-END SDLC", 9, "91D5D0", bold=True)
    footer(slide, 10, dark=True)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build_deck()